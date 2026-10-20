"""
Проверка логики v3 без интернета: обращения, статусы, анонимные ответы, оценки,
напоминание 48 ч, сводка, выгрузка Excel, анонсы с явкой, опросы, новости,
ссылки/контакты/документы, статистика и права доступа API.

    python3 tests/test_core.py
"""
from __future__ import annotations

import asyncio
import datetime as dt
import hashlib
import hmac
import io
import json
import os
import sys
import tempfile
import time
import zipfile
from pathlib import Path
from urllib.parse import quote

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))
TMP = Path(tempfile.mkdtemp())
os.environ.update(BOT_TOKEN="123456:TEST-TOKEN", ADMIN_IDS="1,2", SUPER_CHAT_ID="-100500", TOPIC_ID="7",
                  WEBAPP_URL="https://app.example", DATA_DIR=str(TMP / "data"))

import stubs  # noqa: E402

stubs.install()
import config  # noqa: E402

config.DOCS_DIR = TMP / "docs"
config.DOCS_DIR.mkdir(parents=True)
config.MEDIA_DIR = TMP / "data" / "media"
config.NOTIFY_ADMINS_DM = False
import store  # noqa: E402

store.reset_for_tests(TMP / "data" / "nab.db")
import api  # noqa: E402
import content  # noqa: E402
import reports  # noqa: E402
import stats  # noqa: E402
import xlsx  # noqa: E402
from stubs import FakeBot, kb_buttons  # noqa: E402

ok = 0


def check(cond, msg):
    global ok
    if not cond:
        raise AssertionError(msg)
    ok += 1


def init_data(uid, first="Аня", username="anya"):
    user = json.dumps({"id": uid, "first_name": first, "username": username}, ensure_ascii=False)
    fields = {"auth_date": str(int(time.time())), "query_id": "Q", "user": user}
    check_str = "\n".join(f"{k}={v}" for k, v in sorted(fields.items()))
    sec = hmac.new(b"WebAppData", config.TOKEN.encode(), hashlib.sha256).digest()
    fields["hash"] = hmac.new(sec, check_str.encode(), hashlib.sha256).hexdigest()
    return "&".join(f"{k}={quote(v)}" for k, v in fields.items())


STUDENT = {"id": 777, "first_name": "Аня", "username": "anya"}
OTHER = {"id": 888, "first_name": "Петя", "username": ""}


async def run():
    bot = FakeBot()

    # ---------- 1. обращение с профилем: карточка в группе с кнопками ----------
    good, rid, msg = await reports.create(bot, "question", STUDENT, "Когда пересдача по матану?", source="app")
    check(good and rid, f"обращение создано: {msg}")
    sm = bot.of("send_message")[-1]
    check(sm["chat_id"] == -100500 and sm["message_thread_id"] == 7, "ушло в рабочий чат и тему")
    check("№%d" % rid in sm["text"] and "@anya" in sm["text"], "в карточке номер и автор")
    datas = [b.get("callback_data") for b in kb_buttons(sm["reply_markup"])]
    check(f"rp:take:{rid}" in datas and f"rp:done:{rid}" in datas and f"rp:ask:{rid}" in datas, "кнопки взять/решено/ответить")
    r = reports.get(rid)
    check(r["status"] == "new" and r["user_id"] == 777 and r["message_id"], "статус new, сообщение сохранено")

    # ---------- 2. анонимное обращение: в базе нет ID, но ответ доходит ----------
    good, arid, _ = await reports.create(bot, "panic", STUDENT, "Мне страшно говорить открыто", anon=True)
    ar = reports.get(arid)
    check(ar["user_id"] is None and "777" not in (ar["anon_ref"] or "") and ar["user_name"].startswith("Анонимно"),
          "анонимное: ID и имя не хранятся открыто")
    check("@anya" not in bot.of("send_message")[-1]["text"], "в группе нет имени анонима")
    check(reports.author_of(ar) == 777, "но бот может вычислить чат автора для ответа")
    check(reports.is_author(ar, 777) and not reports.is_author(ar, 888), "автор узнаётся только по своей подписи")
    check(any(x["id"] == arid for x in map(dict, reports.mine(777))), "анонимное видно автору в «Мои обращения»")
    check(not any(x["id"] == arid for x in map(dict, reports.mine(888))), "и не видно другим")

    # ---------- 3. ответ руководства анонимному автору ----------
    n0 = len(bot.calls)
    good, err = await reports.team_reply(bot, arid, "Мы рядом, напиши подробнее", "Алина")
    check(good, f"ответ доставлен: {err}")
    to_author = [k for n, k in bot.calls[n0:] if n == "send_message" and k["chat_id"] == 777]
    check(to_author and "Мы рядом" in to_author[0]["text"] and "Алина" not in to_author[0]["text"], "автор получил ответ без имени админа")
    check(reports.get(arid)["status"] == "work", "после ответа статус «в работе»")
    check(reports.unread(777) == 1, "у автора 1 непрочитанный ответ")
    reports.mark_seen(777)
    check(reports.unread(777) == 0, "после просмотра — 0")

    # ---------- 4. автор отвечает → в группу ответом на карточку ----------
    good, err = await reports.user_reply(bot, arid, 777, "Спасибо, всё расскажу")
    gm = bot.of("send_message")[-1]
    check(good and gm["chat_id"] == "-100500" or gm["chat_id"] == -100500, "ответ автора ушёл в группу")
    check("анонимно" in gm["text"] and gm["reply_parameters"]["message_id"] == ar["message_id"], "ответом на исходную карточку")
    good, err = await reports.user_reply(bot, arid, 888, "чужой")
    check(not good, "чужой не может писать по чужому обращению")

    # ---------- 5. статусы: взять, решено → автору звёзды ----------
    _, e1 = await reports.set_status(bot, rid, "take", "Мусо")
    check(not e1 and reports.get(rid)["taken_by"] == "Мусо", "взято в работу")
    _, e2 = await reports.set_status(bot, rid, "take", "Алина")
    check("Уже в работе" in e2, "повторно взять нельзя")
    edits = bot.of("edit_message_reply_markup")
    check(edits and "Мусо" in kb_buttons(edits[-1]["reply_markup"])[0]["text"], "кнопка в группе показывает, кто взял")
    n0 = len(bot.calls)
    _, e3 = await reports.set_status(bot, rid, "done", "Мусо")
    rate_msg = [k for n, k in bot.calls[n0:] if n == "send_message" and k["chat_id"] == 777][-1]
    stars_btn = [b["callback_data"] for b in kb_buttons(rate_msg["reply_markup"])]
    check(stars_btn == [f"rate:{rid}:{i}" for i in range(1, 6)], "автору пришли 5 звёзд")

    # ---------- 6. оценка и отзыв ----------
    good, err = await reports.rate(bot, rid, 888, 5)
    check(not good, "оценить может только автор")
    good, err = await reports.rate(bot, rid, 777, 4)
    check(good and reports.get(rid)["rating"] == 4, "оценка сохранена")
    good, err = await reports.rate(bot, rid, 777, None, "Быстро ответили, спасибо!")
    check(reports.get(rid)["rating_text"] == "Быстро ответили, спасибо!", "отзыв сохранён")
    check("★★★★☆" in bot.of("send_message")[-1]["text"], "оценка звёздами ушла в группу")
    good, err = await reports.rate(bot, arid, 777, 5)
    check(not good, "нельзя оценить нерешённое")
    _, e4 = await reports.set_status(bot, rid, "reopen", "Алина")
    check(not e4 and reports.get(rid)["status"] == "work", "вернули в работу")
    check(reports.get(rid)["rating"] is None, "после возврата в работу старая оценка сброшена")
    await reports.set_status(bot, rid, "done", "Алина")
    await reports.rate(bot, rid, 777, 4)

    # ---------- 7. напоминание о просроченных (48 ч) ----------
    good, old_id, _ = await reports.create(bot, "initiative", OTHER, "Сделать кофе-поинт")
    store.x("UPDATE reports SET created_at=? WHERE id=?", (int(time.time()) - 50 * 3600, old_id))
    n = await reports.sla_tick(bot)
    check(n == 1 and "ждёт уже 2 дня" in bot.of("send_message")[-1]["text"], "напоминание в группе через 48 ч")
    check(await reports.sla_tick(bot) == 0, "повторно — не раньше чем через сутки")
    check(await reports.sla_tick(bot, time.time() + 25 * 3600) == 1, "через сутки — ещё раз")
    store.set_setting("sla_hours", 72)
    check(len(reports.overdue()) == 0, "срок меняется в настройках (72 ч)")
    store.set_setting("sla_hours", 48)

    # ---------- 8. сводка и выгрузка Excel ----------
    d = reports.digest_text()
    check("Сводка обращений" in d and "Новых:" in d and "№" in d, "сводка собирается")
    y, m = dt.date.today().year, dt.date.today().month
    raw, fname, cap = stats.export_month(y, m)
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    check(wb.sheetnames == ["Обращения", "Сводка", "Анонсы и явка", "Опросы"], f"листы Excel: {wb.sheetnames}")
    ws = wb["Обращения"]
    check(ws.max_row >= 4 and ws["A1"].value == "№" and ws["G1"].value == "Автор", "строки обращений с шапкой")
    authors = [ws.cell(row=i, column=7).value for i in range(2, ws.max_row + 1)]
    check("Анонимно 🕵️‍♂️" not in authors and None in authors or "" in authors, "у анонимных автор пустой")
    check(any(ws.cell(row=i, column=16).value == 4 for i in range(2, ws.max_row + 1)), "оценка в таблице")
    check(fname.endswith(".xlsx") and "Отчёт" in cap, "имя файла и подпись")
    sent = await stats.send_export(bot, [1, 2], y, m)
    check(sent == 2 and bot.of("send_document")[-1]["document"].kwargs["filename"] == fname, "отчёт ушёл двум админам")

    # ---------- 9. анонсы: явка, напоминания ----------
    tomorrow = (content.now_msk() + dt.timedelta(days=1)).date().isoformat()
    eid, err = content.event_save({"title": "Выезд актива", "date": tomorrow, "time": "18:00", "place": "Масловка",
                                   "rsvp": True, "kind": "trip"}, 1)
    check(eid and not err, "анонс создан")
    check(content.set_rsvp(eid, 777, "yes", "Аня (@anya)")[0], "Аня придёт")
    check(content.set_rsvp(eid, 888, "no", "Петя")[0], "Петя не придёт")
    v = content.event_view(store.q("SELECT * FROM events WHERE id=?", (eid,), one=True), 777)
    check(v["counts"] == {"yes": 1, "no": 1} and v["my"] == "yes" and v["subscribed"], "счётчики явки и подписка идущего")
    check(content.attendees(eid)["yes"] == ["Аня (@anya)"], "руководство видит, кто придёт")
    kb = kb_buttons(content.event_kb(store.q("SELECT * FROM events WHERE id=?", (eid,), one=True), 777))
    check(kb[0]["text"].startswith("✅") and kb[0]["callback_data"] == f"ev:yes:{eid}", "кнопки «Приду / Не приду» в чате")
    at = content.event_dt(store.q("SELECT * FROM events WHERE id=?", (eid,), one=True))
    n = await content.event_tick(bot, at - dt.timedelta(hours=20))
    check(n == 1 and "Напоминание на завтра" in bot.of("send_message")[-1]["text"], "напоминание идущим за сутки")
    check(await content.event_tick(bot, at - dt.timedelta(hours=19)) == 0, "второй раз не шлём")
    check(await content.event_tick(bot, at - dt.timedelta(hours=1)) == 1, "за 2 часа — «уже скоро»")
    eid2, _ = content.event_save({"title": "Дедлайн анкеты", "date": tomorrow, "rsvp": False}, 1)
    check(not content.set_rsvp(eid2, 777, "yes")[0], "без явки — отметка не нужна")
    sent, total = await content.announce(bot, eid)
    check(total >= 2 and sent == total, "анонс разослан всем пользователям")

    # ---------- 10. опросы ----------
    pid, err = content.poll_save({"question": "Куда поедем на выезд?", "options": ["Подмосковье", "Питер", "Казань"]}, 1)
    check(pid and not err, "опрос создан")
    v = content.poll_view(store.q("SELECT * FROM polls WHERE id=?", (pid,), one=True), 777)
    check(v["votes"] is None, "до голоса результаты скрыты")
    v, err = content.poll_vote(pid, 777, [1])
    check(v and v["mine"] == [1] and v["votes"] == [0, 1, 0], "голос учтён")
    v, _ = content.poll_vote(pid, 777, [2])
    check(v["votes"] == [0, 0, 1], "переголосовать можно, двойного голоса нет")
    check("Питер" in content.poll_text(v) and "▰" in content.poll_bars(v), "результаты в чате полосками")
    err = content.poll_save({"id": pid, "question": "Куда?", "options": ["А", "Б"]}, 1)[1]
    check("нельзя менять" in err, "варианты нельзя менять после голосов")
    store.x("UPDATE polls SET closed=1 WHERE id=?", (pid,))
    check(content.poll_vote(pid, 888, [0])[0] is None, "в закрытом опросе голосовать нельзя")

    # ---------- 11. новости, ссылки, контакты, документы ----------
    nid, err = content.news_save({"title": "Набор в актив", "body": "Ждём тебя!", "btnText": "Анкета", "btnUrl": "forms.gle/x", "pinned": True}, "Алина")
    check(nid and content.news_list()[0]["btnUrl"] == "https://forms.gle/x", "новость с кнопкой")
    check(content.news_save({"title": "x"}, "А")[0] is None, "слишком короткий заголовок")
    ok_n, tot = await content.news_broadcast(bot, nid)
    check(ok_n == tot and "Набор в актив" in bot.of("send_message")[-1]["text"], "новость разослана")
    check(content.links_save([{"t": "Канал", "u": "@NABLABFU", "k": "Telegram"}])[0] and content.links()[0]["u"] == "https://t.me/NABLABFU",
          "ссылки редактируются, @ превращается в t.me")
    check(not content.links_save([{"t": "", "u": ""}])[0], "пустая ссылка не сохраняется")
    check(content.people_save([{"name": "Алина Зотова", "role": "Председатель", "tg": "@jannalli"}])[0]
          and content.people()[0]["ini"] == "АЗ" and content.people()[0]["tg"] == "jannalli", "контакты: инициалы и ник")
    did, err = content.doc_upload(b"%PDF-1.4 test", "Положение о стипендии", "Как получить")
    check(did == "10" and (config.DOCS_DIR / "10. Положение о стипендии.pdf").exists(), "PDF загружен как №10")
    check(any(d["id"] == "10" and d["custom"] for d in content.docs()), "документ появился в списке")
    check(content.doc_upload(b"not a pdf", "x")[0] is None, "не-PDF отклоняется")
    content.doc_meta_save("3", hidden=True)
    check(not any(d["id"] == "3" for d in content.docs()) and any(d["id"] == "3" for d in content.docs(True)), "стандартный документ можно скрыть")
    check(content.doc_delete("10")[0] and not any(d["id"] == "10" for d in content.docs(True)), "загруженный — удалить")
    reg = json.loads((config.DATA_DIR / "docs.json").read_text(encoding="utf-8"))
    check(any(x["id"] == "1" for x in reg), "реестр для сборки страниц обновляется")

    # ---------- 12. статистика анонимная ----------
    store.seen(777)
    store.seen(777)
    store.seen(888)
    store.hit("screen", "faq")
    store.hit("search", "стипендия")
    dsh = stats.dashboard(30)
    check(dsh["active"]["today"] == 2 and dsh["active"]["month"] == 2, "уникальные за день и месяц")
    check(store.count("SELECT COUNT(*) FROM stat_uniq WHERE uh LIKE '%777%'") == 0, "в статистике нет Telegram ID")
    check(dsh["reports"]["total"] >= 3 and dsh["reports"]["avgRating"] == 4.0, "метрики обращений")
    check({"key": "faq", "n": 1} in dsh["screens"], "топ экранов")

    # ---------- 13. API: подпись Telegram и права ----------
    ctx = api.Ctx(bot=bot)
    st, data = await api.dispatch(ctx, "/api/bootstrap", "garbage", {}, {})
    check(st == 401, "без подписи — 401")
    st, data = await api.dispatch(ctx, "/api/bootstrap", init_data(777), {}, {})
    check(st == 200 and not data["data"]["me"]["isAdmin"] and data["data"]["news"], "bootstrap студента")
    st, data = await api.dispatch(ctx, "/api/admin/summary", init_data(777), {}, {})
    check(st == 403, "студенту Пульт закрыт")
    st, data = await api.dispatch(ctx, "/api/admin/summary", init_data(1, "Алина", "jannalli"), {}, {})
    check(st == 200 and "users" in data["data"], "админу Пульт открыт")
    st, data = await api.dispatch(ctx, "/api/report", init_data(999, "Новый", "new"),
                                  {"kind": "repair", "text": "Сломан стул", "location": "ауд. 412", "anon": "1"},
                                  {"photos": [b"\xff\xd8\xffJPEG1", b"\xff\xd8\xffJPEG2"]})
    check(st == 200 and data["id"], "обращение из приложения с фото")
    check(bot.of("send_media_group")[-1]["media"][0].kwargs["caption"].startswith("📎"), "фото ушли альбомом")
    st, data = await api.dispatch(ctx, "/api/report", init_data(999, "Новый", "new"), {"kind": "repair", "text": "ещё"}, {})
    check(st == 429, "слишком часто — 429")
    st, data = await api.dispatch(ctx, "/api/my-reports", init_data(999, "Новый", "new"), {}, {})
    check(st == 200 and data["data"][0]["anon"] and data["data"][0]["location"] == "ауд. 412", "мои обращения")
    st, data = await api.dispatch(ctx, "/api/admin/report-action", init_data(2, "Мусо", ""),
                                  {"id": data["data"][0]["id"], "action": "reply", "text": "Починим завтра"}, {})
    check(st == 200 and data["report"]["status"] == "work", "ответ из Пульта")
    st, data = await api.dispatch(ctx, "/api/admin/reports", init_data(2), {"filter": "open"}, {})
    check(st == 200 and data["counts"]["open"] >= 2, "список обращений в Пульте")
    st, data = await api.dispatch(ctx, "/api/event-rsvp", init_data(999), {"id": eid, "answer": "yes"}, {})
    check(st == 200 and data["event"]["counts"]["yes"] == 2, "явка из приложения")
    st, data = await api.dispatch(ctx, "/api/admin/settings", init_data(1), {"slaHours": 24, "digestOn": False}, {})
    check(st == 200 and data["settings"]["slaHours"] == 24 and not data["settings"]["digestOn"], "настройки Пульта")
    st, data = await api.dispatch(ctx, "/api/track", init_data(999), {"events": [{"n": "screen", "k": "map"}, {"n": "hack", "k": "x"}]}, {})
    check(st == 200 and store.count("SELECT COUNT(*) FROM stat WHERE name='hack'") == 0, "в статистику попадают только разрешённые события")
    st, data = await api.dispatch(ctx, "/api/admin/news-save", init_data(1), {"title": "Тест", "body": "x", "send": False}, {})
    check(st == 200 and data["id"], "новость из Пульта")
    st, data = await api.dispatch(ctx, "/api/admin/upload-image", init_data(1), {}, {"file": [b"\x89PNG\r\n\x1a\n" + b"0" * 20 + b"IEND\xaeB`\x82"]})
    check(st == 200 and data["name"].endswith(".png") and (config.MEDIA_DIR / data["name"]).exists(), "картинка новости")
    st, data = await api.dispatch(ctx, "/api/admin/upload-image", init_data(1), {}, {"file": [b"<script>"]})
    check(st == 400, "не-картинка отклоняется")

    # ---------- 13б. безопасность: пути, анонимность, лимиты ----------
    outside = TMP / "outside"
    outside.mkdir()
    (outside / "secret.leak.pdf").write_bytes(b"%PDF secret")
    for bad in ("../outside/secret", "*", "../../**/*", "1/../../outside/secret", ""):
        check(content.doc_files(bad) == [], f"путь «{bad}» не открывает файлы")
        st, _ = await api.dispatch(ctx, "/api/doc", init_data(777), {"id": bad}, {})
        check(st == 404, f"/api/doc «{bad}» → 404")
    check(not content.doc_delete("../outside/secret")[0] and (outside / "secret.leak.pdf").exists(), "удалить чужой файл нельзя")
    eid3, _ = content.event_save({"title": "Явка", "date": tomorrow, "time": "20:00", "rsvp": True}, 1)
    content.set_rsvp(eid3, 555, "yes", "x")
    content.set_rsvp(eid3, 555, "no", "x")
    check(555 not in content._recipients(store.q("SELECT * FROM events WHERE id=?", (eid3,), one=True)), "«не приду» — без напоминаний")
    good, an_id, _ = await reports.create(bot, "question", {"id": 4242424242, "first_name": "Скрытный", "username": "hidden_user"}, "тайна", anon=True)
    raw_db = b"".join(p.read_bytes() for p in (TMP / "data").glob("nab.db*"))
    check(b"4242424242" not in raw_db and b"hidden_user" not in raw_db, "ID и ник анонима нигде в базе")
    good, an2, _ = await reports.create(bot, "question", {"id": 4242424242, "first_name": "Скрытный"}, "ещё тайна", anon=True)
    check(reports.get(an_id)["anon_tag"] != reports.get(an2)["anon_tag"], "анонимные обращения одного человека не связаны между собой")
    check({an_id, an2} <= {r["id"] for r in reports.mine(4242424242)}, "но автор видит оба у себя")
    flood = [await reports.user_reply(bot, an_id, 4242424242, f"флуд {i}") for i in range(12)]
    check(sum(1 for g, _ in flood if g) == 8, "ответы автора ограничены — группу не завалить")
    # N1: именное и анонимное обращение одного человека не связаны через карту сообщений
    good, named, _ = await reports.create(bot, "question", {"id": 4242424242, "first_name": "Скрытный"}, "открыто", anon=False)
    reports.map_msg(4242424242, 9001, named, "private")
    reports.map_msg(4242424242, 9002, an_id, "private")
    keys = [r["chat_key"] for r in store.q("SELECT chat_key FROM msg_map WHERE message_id IN (9001, 9002)")]
    check(len(set(keys)) == 2, "у каждого сообщения в личке свой ключ — анонимное не связать с именным")
    check(reports.lookup(4242424242, 9002)["report_id"] == an_id and reports.lookup(4242424243, 9002) is None,
          "ответ в личке по-прежнему находит своё обращение")
    # N3: звёзды и отзыв — разные лимиты
    store.x("UPDATE reports SET status='done' WHERE id=?", (an2,))
    taps = [await reports.rate(bot, an2, 4242424242, value=v) for v in (5, 4, 5, 4, 5, 4, 1, 3)]
    rv = await reports.rate(bot, an2, 4242424242, review="Спасибо, помогли")
    check(all(g for g, _ in taps) and rv[0], "частые тапы по звёздам не блокируют отзыв словами")
    # N6: тип опроса не меняется после голосов
    pid, _ = content.poll_save({"question": "Тип?", "options": ["а", "б"], "multi": True}, 1)
    content.poll_vote(pid, 31337, [0, 1])
    _, err = content.poll_save({"id": pid, "question": "Тип?", "options": ["а", "б"], "multi": False}, 1)
    check(bool(err), "одиночный/множественный выбор нельзя переключить после голосов")
    # #4: подпись к фото считается в UTF-16 (эмодзи = 2)
    cap = content.news_text({"title": "🎉" * 140, "body": "🔥" * 3000, "image": "x"})
    import re as _re
    check(len(_re.sub("<[^>]+>", "", cap).encode("utf-16-le")) // 2 <= 1024, "подпись с эмодзи укладывается в 1024")

    # ---------- 14. xlsx: экранирование и спецсимволы ----------
    raw = xlsx.build([("Тест", [["a", "b"], ["<&>\x01", 5]], None)])
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    check(wb["Тест"]["A2"].value == "<&>" and wb["Тест"]["B2"].value == 5, "спецсимволы и числа в xlsx")

    print(f"✅ core: {ok} проверок пройдено")


if __name__ == "__main__":
    asyncio.run(run())
